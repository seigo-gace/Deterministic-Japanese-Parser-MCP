#!/usr/bin/env python3
"""Build-time acquisition of freely reusable Japanese semantic resources.

No LLM calls, no dictionary-site scraping, no invented meanings. Every emitted
record has source-derived meaning evidence plus source/license provenance.
"""
from __future__ import annotations
import argparse, gzip, hashlib, json, shutil, sqlite3, tempfile, unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Any
from urllib.request import Request, urlopen

UA="Deterministic-Japanese-Parser-MCP/public-semantic-source-acquisition"
PLACEHOLDERS={"","pending review","meaning pending review","unknown meaning","tbd","意味確認待ち","意味・機能はevidence確認待ち"}

def norm(v:Any)->str: return unicodedata.normalize("NFKC",str(v or "")).strip()
def line(v:Any)->str: return json.dumps(v,ensure_ascii=False,sort_keys=True,separators=(",",":"))
def sha(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""): h.update(b)
    return h.hexdigest()
def meaning_ok(vals:list[str])->bool: return any(norm(v).casefold() not in PLACEHOLDERS for v in vals)
def download(url:str,target:Path)->str:
    target.parent.mkdir(parents=True,exist_ok=True)
    with urlopen(Request(url,headers={"User-Agent":UA}),timeout=180) as r,target.open("wb") as f:
        shutil.copyfileobj(r,f,1024*1024)
    return sha(target)
def get_json(url:str)->dict[str,Any]:
    with urlopen(Request(url,headers={"User-Agent":UA,"Accept":"application/vnd.github+json"}),timeout=120) as r:
        obj=json.load(r)
    if not isinstance(obj,dict): raise RuntimeError(f"expected object: {url}")
    return obj
def release_asset(spec:dict[str,Any],target:Path)->tuple[str,str,str]:
    release=get_json(f"https://api.github.com/repos/{spec['repository']}/releases/tags/{spec['release_tag']}")
    xs=[a for a in release.get("assets",[]) if a.get("name")==spec["asset_name"]]
    if len(xs)!=1: raise RuntimeError(f"release asset not unique: {spec['id']}")
    a=xs[0]; expected=str(a.get("digest") or "")
    expected=expected[7:] if expected.startswith("sha256:") else ""
    actual=download(str(a["browser_download_url"]),target)
    if expected and expected!=actual: raise RuntimeError(f"source digest mismatch: {spec['id']}")
    return actual,expected,str(a["browser_download_url"])
def cols(c:sqlite3.Connection,t:str)->list[str]: return [str(r[1]) for r in c.execute(f"PRAGMA table_info({t})")]
def pick(cs:list[str],*names:str,required:bool=True)->str|None:
    m={c.casefold():c for c in cs}
    for n in names:
        if n.casefold() in m:return m[n.casefold()]
    if required: raise RuntimeError(f"missing column {names}; available={cs}")
    return None

def normalize_wnja(spec:dict[str,Any],asset:Path,source_sha:str,out:Path)->dict[str,Any]:
    out.parent.mkdir(parents=True,exist_ok=True); db=out.parent/"wnjpn.db"
    with gzip.open(asset,"rb") as s,db.open("wb") as d: shutil.copyfileobj(s,d,1024*1024)
    c=sqlite3.connect(str(db))
    try:
        tables={str(r[0]) for r in c.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if not {"word","sense","synset_def"}<=tables: raise RuntimeError("WNJA required tables missing")
        wc,sc,dc=cols(c,"word"),cols(c,"sense"),cols(c,"synset_def")
        wid,lemma=pick(wc,"wordid"),pick(wc,"lemma")
        pron,pos=pick(wc,"pron",required=False),pick(wc,"pos",required=False)
        sw,ss,sl=pick(sc,"wordid"),pick(sc,"synset"),pick(sc,"lang",required=False)
        ds,dl,dt=pick(dc,"synset"),pick(dc,"lang",required=False),pick(dc,"def","definition")
        defs:dict[str,list[str]]=defaultdict(list)
        q=f'SELECT "{ds}","{dt}"'+(f',"{dl}"' if dl else '')+' FROM synset_def'
        for r in c.execute(q):
            language=norm(r[2]).casefold() if dl else "jpn"
            text=norm(r[1])
            if text and language in {"jpn","ja"}: defs[str(r[0])].append(text)
        ex:dict[str,list[str]]=defaultdict(list)
        if "synset_ex" in tables:
            ec=cols(c,"synset_ex"); es,et,el=pick(ec,"synset"),pick(ec,"def","example","ex"),pick(ec,"lang",required=False)
            q=f'SELECT "{es}","{et}"'+(f',"{el}"' if el else '')+' FROM synset_ex'
            for r in c.execute(q):
                language=norm(r[2]).casefold() if el else "jpn"; text=norm(r[1])
                if text and language in {"jpn","ja"}: ex[str(r[0])].append(text)
        rel:dict[str,list[dict[str,str]]]=defaultdict(list)
        if "synlink" in tables:
            lc=cols(c,"synlink"); a,b,l=pick(lc,"synset1"),pick(lc,"synset2"),pick(lc,"link","relation")
            for r in c.execute(f'SELECT "{a}","{b}","{l}" FROM synlink'):
                rel[str(r[0])].append({"type":norm(r[2]),"target":norm(r[1])})
        fields=[f'w."{wid}"',f'w."{lemma}"',f's."{ss}"',f'w."{pron}"' if pron else 'NULL',f'w."{pos}"' if pos else 'NULL']
        where=f' WHERE lower(s."{sl}") IN (\'jpn\',\'ja\')' if sl else ''
        rows=list(c.execute(f'SELECT {",".join(fields)} FROM word w JOIN sense s ON w."{wid}"=s."{sw}"{where}'))
        syns:dict[str,list[str]]=defaultdict(list)
        for _,w,s,_,_ in rows:
            w=norm(w)
            if w and w not in syns[str(s)]: syns[str(s)].append(w)
        missing=[]; records=[]
        for wordid,w,s,p,part in rows:
            surface=norm(w); synset=str(s); meanings=sorted(set(defs.get(synset,[])))
            if not surface: continue
            if not meaning_ok(meanings): missing.append({"surface":surface,"synset":synset}); continue
            records.append({
              "id":f"WNJA-{synset}-{wordid}","surface":surface,"readings":[norm(p)] if norm(p) else [],
              "part_of_speech":[norm(part)] if norm(part) else [synset.rsplit('-',1)[-1]],"meanings":meanings,"label":meanings[0],
              "domains":["general","wordnet"],"dataset":spec["title"],"license":spec["license"],"source_id":f"{synset}:{wordid}",
              "source_url":spec["homepage"],"source_sha256":source_sha,"attribution":spec.get("attribution",""),
              "parameters":{"synset":synset,"synonyms":sorted(syns.get(synset,[])),"relations":sorted(rel.get(synset,[]),key=lambda x:(x['type'],x['target'])),"examples":sorted(set(ex.get(synset,[])))}
            })
        if missing: raise RuntimeError(f"WNJA meaning evidence missing: count={len(missing)} sample={missing[:10]}")
        records.sort(key=lambda x:x["id"])
        with out.open("w",encoding="utf-8",newline="\n") as f:
            for r in records:f.write(line(r)+"\n")
        return {"records":len(records),"missing_meaning_records":0,"normalized_sha256":sha(out),"source_tables":sorted(tables)}
    finally:
        c.close(); db.unlink(missing_ok=True)

def h(v:Any)->str:return "".join(norm(v).replace("\n"," ").split()).casefold()
def normalize_okinawa(spec:dict[str,Any],asset:Path,source_sha:str,out:Path)->dict[str,Any]:
    import openpyxl
    wb=openpyxl.load_workbook(asset,read_only=True,data_only=True); records=[]; missing=[]; parsed=0
    try:
        for ws in wb.worksheets:
            header=None; hm={}
            for rn,row in enumerate(ws.iter_rows(values_only=True),1):
                vv=[norm(v) for v in row]; nn=[h(v) for v in vv]
                si=next((i for i,v in enumerate(nn) if "沖縄語" in v or v in {"見出し語","語形"}),None)
                mi=next((i for i,v in enumerate(nn) if "意味" in v),None)
                if si is not None and mi is not None: header=rn; hm={v:i for i,v in enumerate(nn) if v}; break
                if rn>=30:break
            if header is None: continue
            parsed+=1
            def ix(*ns:str)->int|None:
                return next((i for k,i in hm.items() if any(n in k for n in ns)),None)
            si,mi,pi,ai=ix("沖縄語","見出し語","語形"),ix("意味"),ix("品詞"),ix("アクセント")
            if si is None or mi is None: raise RuntimeError(f"Okinawa required columns missing: {ws.title}")
            for rn,row in enumerate(ws.iter_rows(min_row=header+1,values_only=True),header+1):
                vv=list(row); surface=norm(vv[si]) if si<len(vv) else ""
                if not surface: continue
                meaning=norm(vv[mi]) if mi<len(vv) else ""
                if not meaning_ok([meaning]): missing.append({"sheet":ws.title,"row":rn,"surface":surface}); continue
                pos=norm(vv[pi]) if pi is not None and pi<len(vv) else ""; accent=norm(vv[ai]) if ai is not None and ai<len(vv) else ""
                records.append({"id":f"OKINAWA-{ws.title}-{rn:06d}","surface":surface,"meanings":[meaning],"label":meaning,
                  "part_of_speech":[pos] if pos else [],"domains":["dialect","ryukyuan","okinawa-shuri"],"dataset":spec["title"],
                  "license":spec["license"],"source_id":f"{ws.title}:{rn}","source_url":spec["homepage"],"source_sha256":source_sha,
                  "attribution":spec.get("attribution",""),"parameters":{"accent":accent,"sheet":ws.title,"row":rn}})
    finally: wb.close()
    if parsed==0: raise RuntimeError("Okinawa dictionary header not detected")
    if missing: raise RuntimeError(f"Okinawa meaning evidence missing: count={len(missing)} sample={missing[:10]}")
    records.sort(key=lambda x:x["id"]); out.parent.mkdir(parents=True,exist_ok=True)
    with out.open("w",encoding="utf-8",newline="\n") as f:
        for r in records:f.write(line(r)+"\n")
    return {"records":len(records),"missing_meaning_records":0,"normalized_sha256":sha(out)}

def manifest(path:Path)->dict[str,Any]:
    m=json.loads(path.read_text(encoding="utf-8")); src=m.get("sources")
    if not isinstance(src,list): raise RuntimeError("sources must be list")
    ids=[str(x.get("id") or "") for x in src]
    if any(not x for x in ids) or len(ids)!=len(set(ids)): raise RuntimeError("source ids invalid")
    p=m.get("policy") or {}
    for k in ("require_free_access","require_machine_retrievable_data","require_meaning_evidence","require_source_digest","forbid_web_scraping_without_reuse_permission","forbid_placeholder_meanings","fail_closed_on_missing_meaning"):
        if p.get(k) is not True: raise RuntimeError(f"required policy disabled: {k}")
    return m
def validate_ref(path:Path)->dict[str,Any]:
    n=0
    for ln,raw in enumerate(path.read_text(encoding="utf-8").splitlines(),1):
        if not raw.strip():continue
        x=json.loads(raw); n+=1
        if not norm(x.get("surface")) or not meaning_ok([norm(v) for v in x.get("meanings",[])]):raise RuntimeError(f"meaning incomplete {path}:{ln}")
        d=str(x.get("source_sha256") or "")
        if len(d)!=64 or any(c not in "0123456789abcdefABCDEF" for c in d):raise RuntimeError(f"source digest invalid {path}:{ln}")
        if not norm(x.get("license")):raise RuntimeError(f"license missing {path}:{ln}")
    if n==0:raise RuntimeError(f"empty reference: {path}")
    return {"records":n,"sha256":sha(path)}
def acquire(mp:Path,root:Path,rp:Path,lp:Path)->dict[str,Any]:
    m=manifest(mp); raw=root/"raw"; ref=root/"reference"; lic=root/"licenses"; reports=[]; locks=[]
    for s in m["sources"]:
        if s.get("status")!="enabled":continue
        sid=str(s["id"]); adapter=str(s["adapter"]); authority=""
        if adapter=="wnja_sqlite_release":
            a=raw/f"{sid}.db.gz"; source_sha,authority,url=release_asset(s,a); out=ref/f"{sid}.jsonl"; result=normalize_wnja(s,a,source_sha,out)
            license_sha=download(str(s["license_url"]),lic/f"{sid}-LICENSE.txt")
        elif adapter=="okinawa_xlsx":
            a=raw/f"{sid}.xlsx"; url=str(s["url"]); source_sha=download(url,a); out=ref/f"{sid}.jsonl"; result=normalize_okinawa(s,a,source_sha,out); license_sha=None
        else:raise RuntimeError(f"unsupported adapter: {adapter}")
        v=validate_ref(out); reports.append({"source_id":sid,**result,"validation":v})
        locks.append({"source_id":sid,"adapter":adapter,"version":s.get("version") or s.get("release_tag"),"source_url":url,
          "source_sha256":source_sha,"authority_sha256":authority or None,"lock_state":"authoritative-digest" if authority else "computed-digest-needs-canonical-lock",
          "license":s.get("license"),"license_url":s.get("license_url"),"license_sha256":license_sha,"attribution":s.get("attribution",""),
          "normalized_path":str(out.relative_to(root)),"normalized_sha256":v["sha256"],"records":v["records"]})
    if not reports:raise RuntimeError("no enabled sources")
    report={"schema_version":m.get("schema_version"),"enabled_source_count":len(reports),"total_records":sum(x["records"] for x in reports),
      "sources":reports,"meaning_complete":all(x.get("missing_meaning_records")==0 for x in reports),"llm_api_used":False,"web_scraping_used":False}
    rp.parent.mkdir(parents=True,exist_ok=True); lp.parent.mkdir(parents=True,exist_ok=True)
    rp.write_text(json.dumps(report,ensure_ascii=False,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    lp.write_text(json.dumps({"sources":locks},ensure_ascii=False,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    return report
def selftest()->dict[str,Any]:
    with tempfile.TemporaryDirectory() as td:
        r=Path(td); db=r/"t.db"; c=sqlite3.connect(str(db)); c.executescript("""
CREATE TABLE word(wordid INTEGER PRIMARY KEY,lang TEXT,lemma TEXT,pron TEXT,pos TEXT);
CREATE TABLE sense(synset TEXT,wordid INTEGER,lang TEXT,rank INTEGER,lexid INTEGER,freq INTEGER,src TEXT);
CREATE TABLE synset_def(synset TEXT,lang TEXT,def TEXT,sid INTEGER);
CREATE TABLE synset_ex(synset TEXT,lang TEXT,def TEXT,sid INTEGER);
CREATE TABLE synlink(synset1 TEXT,synset2 TEXT,link TEXT,src TEXT);
INSERT INTO word VALUES(1,'jpn','猫','ネコ','n'); INSERT INTO word VALUES(2,'jpn','ねこ','ネコ','n');
INSERT INTO sense VALUES('00000001-n',1,'jpn',1,0,0,'test'); INSERT INTO sense VALUES('00000001-n',2,'jpn',1,0,0,'test');
INSERT INTO synset_def VALUES('00000001-n','jpn','小型の家畜化されたネコ科動物',1);
INSERT INTO synset_ex VALUES('00000001-n','jpn','猫が眠っている',1); INSERT INTO synlink VALUES('00000001-n','00000002-n','hype','test');
"""); c.commit(); c.close(); gz=r/"t.db.gz"
        with db.open("rb") as s,gz.open("wb") as raw:
            with gzip.GzipFile(filename="",mode="wb",fileobj=raw,mtime=0) as d:shutil.copyfileobj(s,d)
        out=r/"o.jsonl"; res=normalize_wnja({"title":"test","license":"test","homepage":"x"},gz,sha(gz),out); v=validate_ref(out)
        if res["records"]!=2 or v["records"]!=2:raise RuntimeError("self-test count")
        return {"status":"PASS","records":2,"sha256":v["sha256"]}
def main()->int:
    p=argparse.ArgumentParser(); sub=p.add_subparsers(dest="cmd",required=True)
    a=sub.add_parser("validate-manifest"); a.add_argument("--manifest",type=Path,required=True)
    sub.add_parser("self-test")
    a=sub.add_parser("acquire"); a.add_argument("--manifest",type=Path,required=True);a.add_argument("--output-root",type=Path,required=True);a.add_argument("--report",type=Path,required=True);a.add_argument("--lock",type=Path,required=True)
    a=p.parse_args()
    if a.cmd=="validate-manifest":
        m=manifest(a.manifest);print(line({"status":"PASS","sources":len(m["sources"]),"enabled":sum(1 for s in m["sources"] if s.get("status")=="enabled")}))
    elif a.cmd=="self-test":print(line(selftest()))
    else:print(line(acquire(a.manifest,a.output_root,a.report,a.lock)))
    return 0
if __name__=="__main__":raise SystemExit(main())
